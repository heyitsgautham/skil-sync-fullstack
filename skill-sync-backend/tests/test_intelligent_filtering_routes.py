"""The /api/filter routes companies use to rank, inspect and export candidates."""

import io
import json

import openpyxl
import pytest

from app.models.application import Application
from app.models.user import User
from app.services.batch_matching_service import BatchMatchingService

API = "/api/filter"


@pytest.fixture
def matched(seeded, db_session):
    """The seeded world with pre-computed matches."""
    BatchMatchingService(db_session).compute_all_matches()
    return seeded


@pytest.fixture
def company(matched, auth_headers):
    return auth_headers(matched.company)


def apply(client, auth_headers, student, internship_id, **kwargs):
    response = client.post(f"/api/internship/{internship_id}/apply", headers=auth_headers(student), **kwargs)
    assert response.status_code == 201, response.text
    return response.json()


# -- ranking applicants ----------------------------------------------------------


def test_applicant_ranking_scores_each_applicant(client, matched, company, auth_headers):
    for name in ("asha", "ravi"):
        apply(client, auth_headers, matched.students[name], matched.backend.id)

    body = client.post(
        f"{API}/rank-candidates/{matched.backend.id}", params={"only_applicants": True}, headers=company
    ).json()

    assert [c["candidate_id"] for c in body["ranked_candidates"]] == [
        matched.students["asha"].id,
        matched.students["ravi"].id,
    ]
    assert body["ranked_candidates"][0]["scoring_breakdown"]["has_tailored"] is False
    assert "applicants" in body["message"]


def test_applicant_ranking_blends_a_tailored_resume(
    client, matched, company, auth_headers, fake_gemini, tmp_path, monkeypatch
):
    monkeypatch.chdir(tmp_path)
    fake_gemini.reply(json.dumps({"skills": {"technical": ["React", "JavaScript", "Python", "FastAPI"], "soft": []}}))
    apply(
        client,
        auth_headers,
        matched.students["ravi"],
        matched.backend.id,
        data={"use_tailored_resume": "true"},
        files={"tailored_resume": ("t.txt", b"Python FastAPI backend services with Docker", "text/plain")},
    )

    body = client.post(
        f"{API}/rank-candidates/{matched.backend.id}", params={"only_applicants": True}, headers=company
    ).json()

    (candidate,) = body["ranked_candidates"]
    assert candidate["scoring_breakdown"]["has_tailored"] is True
    assert "1 with tailored resumes" in body["performance_note"]


def test_applicant_ranking_keeps_an_applicant_with_no_precomputed_match(
    client, matched, company, auth_headers, db_session
):
    """A student who applies before the next batch run has no match row, and used to turn the whole ranking into a 500."""
    from app.models.student_internship_match import StudentInternshipMatch

    apply(client, auth_headers, matched.students["asha"], matched.backend.id)
    late = apply(client, auth_headers, matched.students["ravi"], matched.backend.id)
    db_session.query(StudentInternshipMatch).filter(
        StudentInternshipMatch.student_id == matched.students["ravi"].id
    ).delete()
    db_session.commit()

    response = client.post(
        f"{API}/rank-candidates/{matched.backend.id}", params={"only_applicants": True}, headers=company
    )

    assert response.status_code == 200
    ranked = {c["candidate_id"]: c for c in response.json()["ranked_candidates"]}
    assert set(ranked) == {matched.students["asha"].id, matched.students["ravi"].id}
    ravi = ranked[matched.students["ravi"].id]
    assert ravi["match_score"] == late["match_score"]
    assert ravi["scoring_breakdown"]["final_weight"] == "application score"
    assert ravi["scoring_breakdown"]["base_similarity"] is None


def test_applicant_ranking_without_applicants(client, matched, company):
    body = client.post(
        f"{API}/rank-candidates/{matched.frontend.id}", params={"only_applicants": True}, headers=company
    ).json()
    assert body["ranked_candidates"] == []
    assert body["message"] == "No applicants found for this internship"


def test_discovery_ranking_marks_applicants(client, matched, company, auth_headers):
    apply(client, auth_headers, matched.students["meera"], matched.backend.id)

    body = client.post(f"{API}/rank-candidates/{matched.backend.id}", headers=company).json()

    applied = {c["candidate_id"]: c["has_applied"] for c in body["ranked_candidates"]}
    assert applied[matched.students["meera"].id] is True
    assert applied[matched.students["asha"].id] is False


def test_discovery_ranking_without_precomputed_matches(client, seeded, auth_headers):
    body = client.post(f"{API}/rank-candidates/{seeded.backend.id}", headers=auth_headers(seeded.company)).json()
    assert body["total_candidates"] == 0
    assert "compute-matches" in body["recommendation"]


def test_students_cannot_rank(client, matched, auth_headers):
    headers = auth_headers(matched.students["asha"])
    assert client.post(f"{API}/rank-candidates/{matched.backend.id}", headers=headers).status_code == 403


# -- candidate details -----------------------------------------------------------------


def test_flagged_candidate_details(client, matched, company, db_session):
    ids = [matched.students["asha"].id, matched.students["ravi"].id]
    db_session.query(User).filter(User.id.in_(ids)).update({"phone": "98765 43210"}, synchronize_session=False)
    db_session.commit()

    body = client.get(
        f"{API}/flagged-candidates-details", params={"candidate_ids": f"{ids[0]}, {ids[1]}"}, headers=company
    ).json()

    assert body["total_count"] == 2
    assert all(c["is_flagged"] for c in body["candidates"])
    assert body["candidates"][0]["flag_reason_text"] == "Same Mobile number"


def test_flagged_candidate_details_input_checks(client, matched, company):
    bad = client.get(f"{API}/flagged-candidates-details", params={"candidate_ids": "one,two"}, headers=company)
    assert bad.status_code == 400
    empty = client.get(f"{API}/flagged-candidates-details", params={"candidate_ids": "999999"}, headers=company)
    assert empty.json()["candidates"] == []


def test_match_score_for_one_pair(client, matched, company, fake_gemini, monkeypatch):
    from app.routes import intelligent_filtering

    monkeypatch.setattr(intelligent_filtering.matching_engine, "key_manager", fake_gemini)
    fake_gemini.reply("- Strong Fit")

    body = client.post(
        f"{API}/match-score",
        params={"student_id": matched.students["asha"].id, "internship_id": matched.backend.internship_id},
        headers=company,
    ).json()

    assert body["explanation"] == "- Strong Fit"
    assert body["match_result"]["component_scores"]["skills_match"] == 100.0


def test_compute_matches_modes(client, seeded, auth_headers, db_session):
    headers = auth_headers(seeded.company)

    everything = client.post(f"{API}/compute-matches", headers=headers).json()
    one_student = client.post(
        f"{API}/compute-matches", params={"student_id": seeded.students["ravi"].id}, headers=headers
    ).json()
    one_internship = client.post(
        f"{API}/compute-matches", params={"internship_id": seeded.frontend.id}, headers=headers
    ).json()

    assert everything["statistics"]["matches_computed"] == 6
    assert one_student["statistics"]["matches_computed"] == 2
    assert one_internship["statistics"]["matches_computed"] == 3


# -- filtered, paginated rankings ---------------------------------------------------


def test_filtered_ranking_sorts_and_pages(client, matched, company):
    url = f"{API}/rank-candidates/{matched.backend.id}/filtered"

    first = client.get(url, params={"page_size": 2}, headers=company).json()
    second = client.get(url, params={"page_size": 2, "page": 2}, headers=company).json()
    by_name = client.get(url, params={"sort_by": "name", "sort_order": "asc"}, headers=company).json()

    assert first["total"] == 3 and first["total_pages"] == 2
    assert first["ranked_candidates"][0]["candidate_id"] == matched.students["asha"].id
    assert len(second["ranked_candidates"]) == 1
    assert [c["candidate_name"] for c in by_name["ranked_candidates"]] == ["Asha", "Meera", "Ravi"]


def test_filtered_ranking_by_score_and_applicants(client, matched, company, auth_headers):
    url = f"{API}/rank-candidates/{matched.backend.id}/filtered"
    top = client.get(url, headers=company).json()["ranked_candidates"][0]["match_score"]

    only_top = client.get(url, params={"min_score": int(top)}, headers=company).json()
    assert [c["candidate_id"] for c in only_top["ranked_candidates"]] == [matched.students["asha"].id]

    apply(client, auth_headers, matched.students["ravi"], matched.backend.id)
    applicants = client.get(url, params={"only_applicants": True, "sort_by": "date"}, headers=company).json()
    assert [c["candidate_id"] for c in applicants["ranked_candidates"]] == [matched.students["ravi"].id]


# -- export -------------------------------------------------------------------------


def test_export_csv(client, matched, company):
    response = client.get(f"{API}/export-candidates/{matched.backend.id}", params={"format": "csv"}, headers=company)

    assert response.status_code == 200
    assert response.headers["content-type"].startswith("text/csv")
    lines = response.text.strip().splitlines()
    assert len(lines) == 4
    assert "Asha" in lines[1]


def test_export_xlsx(client, matched, company):
    response = client.get(f"{API}/export-candidates/{matched.backend.id}", params={"format": "xlsx"}, headers=company)

    assert response.status_code == 200
    sheet = openpyxl.load_workbook(io.BytesIO(response.content)).active
    assert sheet.max_row == 7  # title, subtitle, blank row, header, three candidates
    names = {sheet.cell(row=row, column=column).value for row in range(5, 8) for column in range(1, 4)}
    assert {"Asha", "Ravi", "Meera"} <= names


def test_export_rules(client, matched, make_user, auth_headers):
    from app.models.user import UserRole

    other = auth_headers(make_user(role=UserRole.company))
    assert client.get(f"{API}/export-candidates/{matched.backend.id}", headers=other).status_code == 403
    assert client.get(f"{API}/export-candidates/999999", headers=other).status_code == 404


# -- resume parsing --------------------------------------------------------------------


def test_parse_resume_stores_and_indexes_it(
    client, make_user, auth_headers, fake_rag, fake_gemini, tmp_path, monkeypatch, db_session
):
    from app.models.user import UserRole

    monkeypatch.chdir(tmp_path)
    student = make_user(role=UserRole.student)
    fake_gemini.reply(
        json.dumps(
            {
                "skills": {"technical": ["Python", "Django"], "soft": []},
                "experience": [{"start_date": "2023-01", "end_date": "2024-01"}],
                "education": [{"degree": "B.Tech"}],
            }
        )
    )

    response = client.post(
        f"{API}/parse-resume",
        files={"file": ("cv.txt", b"Python Django developer at Acme", "text/plain")},
        headers=auth_headers(student),
    )

    assert response.status_code == 200
    body = response.json()
    assert body["processing_details"]["skills_extracted"] == 2
    assert body["processing_details"]["experience_calculated"] == "1.0 years"
    stored = db_session.query(User).filter(User.id == student.id).one()
    assert sorted(stored.skills) == ["Django", "Python"]
    assert (tmp_path / "app/public/resumes" / f"{student.id}_cv.txt").exists()


def test_parse_resume_rejects_unsupported_files(client, make_user, auth_headers, fake_rag, tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    response = client.post(
        f"{API}/parse-resume",
        files={"file": ("cv.exe", b"x", "application/octet-stream")},
        headers=auth_headers(make_user()),
    )
    assert response.status_code == 400
    assert response.json()["detail"].startswith("Invalid file type")


def test_applications_are_recorded(db_session, client, matched, auth_headers):
    apply(client, auth_headers, matched.students["asha"], matched.backend.id)
    assert db_session.query(Application).count() == 1


def test_candidate_profile_by_public_user_id(client, matched, company, auth_headers):
    student = matched.students["meera"]
    apply(client, auth_headers, student, matched.backend.id)

    response = client.get(f"{API}/candidate-profile/{student.user_id}", headers=company)

    assert response.status_code == 200
    assert response.json()["profile"]["all_skills"] == ["Python", "SQL"]
    assert response.json()["metadata"]["total_applications"] == 1


def test_bulk_parse_stub_is_gone(client, matched, company):
    files = [("files", ("a.txt", b"x", "text/plain"))]
    assert client.post(f"{API}/bulk-parse", files=files, headers=company).status_code == 404
