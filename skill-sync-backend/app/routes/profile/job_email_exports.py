"""CSV and Excel attachments for the job report email."""

import csv
import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.models import Application, Resume, User


def generate_csv_export_from_matches(matches, db, internship):
    """Generate CSV content from matched candidates - matches intelligent_filtering export format"""
    output = io.StringIO()
    writer = csv.writer(output)

    # Write header - same as intelligent_filtering export
    writer.writerow(
        [
            'Candidate Name',
            'Email',
            'Phone',
            'Match Score (%)',
            'Top Matching Skills',
            'Experience (Years)',
            'Education Level',
            'Application Date',
            'Application Status',
            'Key Strengths',
            'Semantic Match (%)',
            'Skills Match (%)',
            'Experience Match (%)',
        ]
    )

    # Sort matches by base_similarity_score descending (highest first)
    sorted_matches = sorted(matches, key=lambda m: m.base_similarity_score or 0, reverse=True)

    # Write data
    for match in sorted_matches:
        student = db.query(User).filter(User.id == match.student_id).first()
        if not student:
            continue

        # Get the active resume for parsed data
        resume = db.query(Resume).filter(Resume.student_id == match.student_id, Resume.is_active == 1).first()

        # Handle case where resume or parsed_data might be None
        parsed_data = {}
        if resume and resume.parsed_data:
            parsed_data = resume.parsed_data
        personal_info = parsed_data.get('personal_info', {})

        # Check if this candidate has applied
        application = (
            db.query(Application)
            .filter(Application.student_id == match.student_id, Application.internship_id == match.internship_id)
            .first()
        )

        score = match.base_similarity_score or 0

        # Get skills from parsed data and match with required skills
        candidate_skills = parsed_data.get('all_skills', [])
        required_skills = (internship or match.internship).required_skills or []
        matched_skills = [s for s in candidate_skills if s.lower() in [rs.lower() for rs in required_skills]]
        top_skills = ', '.join(matched_skills[:5]) if matched_skills else 'N/A'

        # Get education level
        education = parsed_data.get('education', [])
        education_level_str = education[0].get('degree', 'N/A') if education else 'N/A'

        # Get key strengths (from projects and certifications)
        projects = parsed_data.get('projects', [])
        certifications = parsed_data.get('certifications', [])
        key_strengths = []
        if projects:
            key_strengths.append(f"{len(projects)} projects")
        if certifications:
            key_strengths.append(f"{len(certifications)} certifications")
        if matched_skills:
            key_strengths.append(f"{len(matched_skills)}/{len(required_skills)} required skills")
        key_strengths_str = ', '.join(key_strengths) if key_strengths else 'Basic qualification'

        # Get phone from parsed data
        phone_number = personal_info.get('phone', 'N/A')

        # Format phone number for CSV (use formula to force text)
        if phone_number != 'N/A' and phone_number:
            # Ensure phone has proper format with country code
            if not phone_number.startswith('+'):
                phone_number = f"+{phone_number}"
            # Format with space after country code if needed
            if phone_number.startswith('+91-'):
                phone_number = phone_number.replace('+91-', '+91 ')
            # Use Excel formula to force text interpretation
            phone_number = f'="{phone_number}"'

        # Get experience from parsed data
        experience_years = parsed_data.get('total_experience_years', 0)

        # Get application date and status
        app_date = str(application.created_at) if application else 'Not Applied'
        app_status = application.status if application else 'Not Applied'

        writer.writerow(
            [
                student.full_name or 'N/A',
                student.email,
                phone_number,
                f"{score:.2f}",
                top_skills,
                experience_years,
                education_level_str,
                app_date,
                app_status,
                key_strengths_str,
                round(match.semantic_similarity, 2) if match.semantic_similarity is not None else None,
                round(match.skills_match_score, 2) if match.skills_match_score is not None else None,
                round(match.experience_match_score, 2) if match.experience_match_score is not None else None,
            ]
        )

    return output.getvalue()


def generate_excel_export_from_matches(matches, db, internship):
    """Generate Excel (XLSX) content from matched candidates - matches intelligent_filtering export format"""

    # Create workbook and worksheet
    wb = Workbook()
    ws = wb.active
    ws.title = "Candidate Rankings"

    # Define styles
    from openpyxl.styles import Border, Side

    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=12)
    header_alignment = Alignment(horizontal="center", vertical="center")
    border = Border(
        left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin')
    )

    # Add title row (row 1)
    ws.merge_cells('A1:M1')
    ws['A1'] = f"Candidate Rankings - {internship.title if internship else 'All Jobs'}"
    ws['A1'].font = Font(bold=True, size=14)
    ws['A1'].alignment = Alignment(horizontal="center")

    # Add summary row (row 2)
    ws.merge_cells('A2:M2')
    ws['A2'] = f"Exported on: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')} | Total Candidates: {len(matches)}"
    ws['A2'].font = Font(size=10, italic=True)
    ws['A2'].alignment = Alignment(horizontal="center")

    # Define headers - same as intelligent_filtering export
    headers = [
        'Candidate Name',
        'Email',
        'Phone',
        'Match Score (%)',
        'Top Matching Skills',
        'Experience (Years)',
        'Education Level',
        'Application Date',
        'Application Status',
        'Key Strengths',
        'Semantic Match (%)',
        'Skills Match (%)',
        'Experience Match (%)',
    ]

    # Write headers with styling (row 4)
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=4, column=col_num)
        cell.value = header
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment
        cell.border = border

    # Sort matches by base_similarity_score descending (highest first)
    sorted_matches = sorted(matches, key=lambda m: m.base_similarity_score or 0, reverse=True)

    # Write data rows
    row_num = 5
    for match in sorted_matches:
        student = db.query(User).filter(User.id == match.student_id).first()
        if not student:
            continue

        # Get the active resume for parsed data
        resume = db.query(Resume).filter(Resume.student_id == match.student_id, Resume.is_active == 1).first()

        # Handle case where resume or parsed_data might be None
        parsed_data = {}
        if resume and resume.parsed_data:
            parsed_data = resume.parsed_data
        personal_info = parsed_data.get('personal_info', {})

        # Check if this candidate has applied
        application = (
            db.query(Application)
            .filter(Application.student_id == match.student_id, Application.internship_id == match.internship_id)
            .first()
        )

        score = match.base_similarity_score or 0

        # Get skills from parsed data and match with required skills
        candidate_skills = parsed_data.get('all_skills', [])
        required_skills = (internship or match.internship).required_skills or []
        matched_skills = [s for s in candidate_skills if s.lower() in [rs.lower() for rs in required_skills]]
        top_skills = ', '.join(matched_skills[:5]) if matched_skills else 'N/A'

        # Get education level
        education = parsed_data.get('education', [])
        education_level_str = education[0].get('degree', 'N/A') if education else 'N/A'

        # Get key strengths (from projects and certifications)
        projects = parsed_data.get('projects', [])
        certifications = parsed_data.get('certifications', [])
        key_strengths = []
        if projects:
            key_strengths.append(f"{len(projects)} projects")
        if certifications:
            key_strengths.append(f"{len(certifications)} certifications")
        if matched_skills:
            key_strengths.append(f"{len(matched_skills)}/{len(required_skills)} required skills")
        key_strengths_str = ', '.join(key_strengths) if key_strengths else 'Basic qualification'

        # Get phone from parsed data
        phone_number = personal_info.get('phone', 'N/A')

        # Format phone number for Excel
        if phone_number != 'N/A' and phone_number:
            # Ensure phone has proper format with country code
            if not phone_number.startswith('+'):
                phone_number = f"+{phone_number}"
            # Format with space after country code if needed
            if phone_number.startswith('+91-'):
                phone_number = phone_number.replace('+91-', '+91 ')

        # Get experience from parsed data
        experience_years = parsed_data.get('total_experience_years', 0)

        # Get application date and status
        app_date = str(application.created_at) if application else 'Not Applied'
        app_status = application.status if application else 'Not Applied'

        # Write row data
        ws.cell(row=row_num, column=1, value=student.full_name or 'N/A')
        ws.cell(row=row_num, column=2, value=student.email)

        # Format phone number and set as text
        phone_cell = ws.cell(row=row_num, column=3, value=phone_number)
        phone_cell.number_format = '@'  # @ means text format in Excel to prevent formula interpretation

        ws.cell(row=row_num, column=4, value=round(score, 2))
        ws.cell(row=row_num, column=5, value=top_skills)
        ws.cell(row=row_num, column=6, value=experience_years)
        ws.cell(row=row_num, column=7, value=education_level_str)
        ws.cell(row=row_num, column=8, value=app_date)
        ws.cell(row=row_num, column=9, value=app_status)
        ws.cell(row=row_num, column=10, value=key_strengths_str)
        ws.cell(
            row=row_num,
            column=11,
            value=round(match.semantic_similarity, 2) if match.semantic_similarity is not None else None,
        )
        ws.cell(
            row=row_num,
            column=12,
            value=round(match.skills_match_score, 2) if match.skills_match_score is not None else None,
        )
        ws.cell(
            row=row_num,
            column=13,
            value=round(match.experience_match_score, 2) if match.experience_match_score is not None else None,
        )

        # Color code only the Match Score cell
        score_cell = ws.cell(row=row_num, column=4)
        if score >= 80:
            score_cell.fill = PatternFill(start_color="C8E6C9", end_color="C8E6C9", fill_type="solid")  # Light green
        elif score >= 60:
            score_cell.fill = PatternFill(start_color="FFE0B2", end_color="FFE0B2", fill_type="solid")  # Light orange
        else:
            score_cell.fill = PatternFill(start_color="FFCDD2", end_color="FFCDD2", fill_type="solid")  # Light red

        row_num += 1

    # Auto-adjust column widths
    for col_num in range(1, len(headers) + 1):
        column_letter = get_column_letter(col_num)
        max_length = 0
        for cell in ws[column_letter]:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except Exception:
                pass
        adjusted_width = min(max_length + 2, 50)  # Cap at 50
        ws.column_dimensions[column_letter].width = adjusted_width

    # Save to bytes
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    return output.getvalue()
