"""GET /api/filter/export-candidates/{internship_id}."""

import io
from datetime import datetime

from fastapi import Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session

from app.database.connection import get_db
from app.models.application import Application
from app.models.internship import Internship
from app.models.resume import Resume
from app.models.user import User
from app.routes.intelligent_filtering.common import router
from app.services.candidate_filters import experience_years, resume_matches
from app.utils.security import get_current_company


def format_phone_for_export(phone, file_format):
    """Format a resume phone number for one export cell, keeping the digits as given.

    CSV wraps the number in ="..." so spreadsheet apps keep it as text instead of dropping a
    leading + or 0. XLSX marks the cell as text in the writer, so the value stays plain.
    """
    if not phone or phone == 'N/A':
        return phone
    # A number without a country code stays as written. Prefixing "+" would turn a local
    # number such as 98765 43210 into a different international one.
    if phone.startswith('+91-'):
        phone = phone.replace('+91-', '+91 ', 1)
    if file_format.lower() == 'xlsx':
        return phone
    return f'="{phone}"'


@router.get("/export-candidates/{internship_id}")
async def export_candidate_rankings(
    internship_id: str,
    format: str = Query("csv", description="Export format: csv or xlsx"),
    # Filtering options (same as filtered endpoint)
    min_score: int | None = Query(None, ge=0, le=100, description="Minimum match score"),
    max_score: int | None = Query(None, ge=0, le=100, description="Maximum match score"),
    skills: str | None = Query(None, description="Comma-separated skills to filter"),
    experience_min: float | None = Query(None, ge=0, description="Minimum years of experience"),
    experience_max: float | None = Query(None, ge=0, description="Maximum years of experience"),
    education_level: str | None = Query(None, description="Education level filter"),
    application_status: str | None = Query(None, description="Filter by application status"),
    only_applicants: bool = Query(False, description="Export only students who applied"),
    # Export options
    export_type: str = Query("filtered", description="Export type: filtered, all, or current_page"),
    page: int = Query(1, ge=1, description="Page number (for current_page export)"),
    page_size: int = Query(10, ge=1, le=100, description="Items per page (for current_page export)"),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_company),
):
    """
    Export candidate rankings to CSV or XLSX format

    Export Formats:
    - **csv**: CSV format (Excel compatible)
    - **xlsx**: Native Excel format with formatting

    Export Types:
    - **filtered**: Export all candidates matching current filters
    - **all**: Export all candidates (ignoring filters)
    - **current_page**: Export only candidates on current page

    Export Fields:
    - Candidate Name, Email, Phone
    - Match Score (%), Top Matching Skills
    - Experience Level, Education Level
    - Application Date, Application Status
    - Key Strengths (brief summary)
    """
    import logging

    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from sqlalchemy import and_

    logger = logging.getLogger(__name__)

    try:
        logger.info(f"📊 Exporting candidates for internship {internship_id} (format: {format})")

        from app.models.student_internship_match import StudentInternshipMatch

        # Get internship
        internship = db.query(Internship).filter(Internship.internship_id == internship_id).first()
        if not internship:
            try:
                int_id = int(internship_id)
                internship = db.query(Internship).filter(Internship.id == int_id).first()
            except ValueError:
                pass

        if not internship:
            raise HTTPException(status_code=404, detail="Internship not found")

        # Verify ownership
        if internship.company_id != current_user.id:
            raise HTTPException(status_code=403, detail="Not authorized to export this data")

        # Build base query
        if only_applicants:
            query = (
                db.query(StudentInternshipMatch, User, Resume, Application)
                .join(User, StudentInternshipMatch.student_id == User.id)
                .join(Resume, and_(Resume.student_id == User.id, Resume.is_active == 1))
                .join(Application, and_(Application.student_id == User.id, Application.internship_id == internship.id))
                .filter(
                    StudentInternshipMatch.internship_id == internship.id,
                    User.is_active == 1,  # Only show active students in rankings
                )
            )
        else:
            query = (
                db.query(StudentInternshipMatch, User, Resume)
                .join(User, StudentInternshipMatch.student_id == User.id)
                .join(Resume, and_(Resume.student_id == User.id, Resume.is_active == 1))
                .filter(
                    StudentInternshipMatch.internship_id == internship.id,
                    User.is_active == 1,  # Only show active students in rankings
                )
            )

        # Apply filters only if not exporting all
        if export_type != "all":
            filters = []

            # Score range filter
            if min_score is not None:
                filters.append(StudentInternshipMatch.base_similarity_score >= min_score)
            if max_score is not None:
                filters.append(StudentInternshipMatch.base_similarity_score <= max_score)

            # Experience filter
            if experience_min is not None or experience_max is not None:
                if experience_min is not None:
                    filters.append(experience_years(Resume.parsed_data) >= experience_min)
                if experience_max is not None:
                    filters.append(experience_years(Resume.parsed_data) <= experience_max)

            # Application status filter
            if only_applicants and application_status:
                filters.append(Application.status == application_status)

            if filters:
                query = query.filter(and_(*filters))

        # Sort by match score (highest first)
        query = query.order_by(StudentInternshipMatch.base_similarity_score.desc())

        # Skills and education live inside JSON lists, so they are checked in Python on the rows
        # the SQL filters already narrowed.
        results = query.all()
        if export_type != "all" and (skills or education_level):
            results = [row for row in results if resume_matches(row[2].parsed_data, skills, education_level)]

        # Apply pagination for current_page export
        if export_type == "current_page":
            offset = (page - 1) * page_size
            results = results[offset : offset + page_size]

        logger.info(f"✅ Found {len(results)} candidates to export")

        # Prepare data for export
        export_data = []
        for row in results:
            match = row[0]
            user = row[1]
            resume = row[2]
            application = row[3] if only_applicants and len(row) > 3 else None

            parsed_data = resume.parsed_data or {}
            personal_info = parsed_data.get('personal_info', {})

            # Get top matching skills
            candidate_skills = parsed_data.get('all_skills', [])
            required_skills = internship.required_skills or []
            matched_skills = [s for s in candidate_skills if s.lower() in [rs.lower() for rs in required_skills]]
            top_skills = ', '.join(matched_skills[:5]) if matched_skills else 'N/A'

            # Get education level
            education = parsed_data.get('education', [])
            education_level_str = education[0].get('degree', 'N/A') if education else 'N/A'

            # Get key strengths (from projects and certifications)
            projects = parsed_data.get('projects', [])
            certifications = parsed_data.get('certifications', [])
            key_strengths = []
            if projects:
                key_strengths.append(f"{len(projects)} projects")
            if certifications:
                key_strengths.append(f"{len(certifications)} certifications")
            if matched_skills:
                key_strengths.append(f"{len(matched_skills)}/{len(required_skills)} required skills")
            key_strengths_str = ', '.join(key_strengths) if key_strengths else 'Basic qualification'

            # Get raw phone number
            phone_number = personal_info.get('phone', 'N/A')

            export_data.append(
                {
                    'Candidate Name': user.full_name,
                    'Email': user.email,
                    'Phone': phone_number,
                    'Match Score (%)': round(match.base_similarity_score, 2),
                    'Top Matching Skills': top_skills,
                    'Experience (Years)': parsed_data.get('total_experience_years', 0),
                    'Education Level': education_level_str,
                    'Application Date': str(application.created_at) if application else 'Not Applied',
                    'Application Status': application.status if application else 'Not Applied',
                    'Key Strengths': key_strengths_str,
                    'Semantic Match (%)': round(match.semantic_similarity, 2)
                    if match.semantic_similarity is not None
                    else None,
                    'Skills Match (%)': round(match.skills_match_score, 2)
                    if match.skills_match_score is not None
                    else None,
                    'Experience Match (%)': round(match.experience_match_score, 2)
                    if match.experience_match_score is not None
                    else None,
                }
            )

        # Generate filename
        internship_title = internship.title.replace(' ', '_').replace('/', '-')
        date_str = datetime.now().strftime('%Y%m%d')
        filename = f"{internship_title}_Candidates_{date_str}"

        for row in export_data:
            row['Phone'] = format_phone_for_export(row.get('Phone'), format)

        # Export based on format
        if format.lower() == 'xlsx':
            # Check if we have data
            if not export_data:
                raise HTTPException(status_code=404, detail="No candidates found to export")

            # Create Excel file with formatting
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Candidate Rankings"

            # Define styles
            header_font = Font(bold=True, color="FFFFFF", size=12)
            header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
            header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            border = Border(
                left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin')
            )

            # Add header with internship info
            ws.merge_cells('A1:M1')
            ws['A1'] = f"Candidate Rankings - {internship.title}"
            ws['A1'].font = Font(bold=True, size=14)
            ws['A1'].alignment = Alignment(horizontal="center")

            ws.merge_cells('A2:M2')
            ws['A2'] = (
                f"Exported on: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')} | Total Candidates: {len(export_data)}"
            )
            ws['A2'].font = Font(size=10, italic=True)
            ws['A2'].alignment = Alignment(horizontal="center")

            # Add column headers (starting from row 4)
            headers = list(export_data[0].keys())
            for col_idx, header in enumerate(headers, start=1):
                cell = ws.cell(row=4, column=col_idx, value=header)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = header_alignment
                cell.border = border

            # Add data rows
            for row_idx, row_data in enumerate(export_data, start=5):
                for col_idx, (key, value) in enumerate(row_data.items(), start=1):
                    cell = ws.cell(row=row_idx, column=col_idx, value=value)
                    cell.border = border

                    # Format phone numbers as text to prevent formula interpretation
                    if key == 'Phone':
                        cell.number_format = '@'  # @ means text format in Excel

                    # Color code only the Match Score cell
                    if key == 'Match Score (%)':
                        if value >= 80:
                            cell.fill = PatternFill(
                                start_color="C8E6C9", end_color="C8E6C9", fill_type="solid"
                            )  # Light green
                        elif value >= 60:
                            cell.fill = PatternFill(
                                start_color="FFE0B2", end_color="FFE0B2", fill_type="solid"
                            )  # Light orange
                        else:
                            cell.fill = PatternFill(
                                start_color="FFCDD2", end_color="FFCDD2", fill_type="solid"
                            )  # Light red

            # Auto-adjust column widths (safer approach)
            for col_idx in range(1, len(headers) + 1):
                max_length = 0
                column_letter = openpyxl.utils.get_column_letter(col_idx)

                # Check header length
                header_cell = ws.cell(row=4, column=col_idx)
                if header_cell.value:
                    max_length = len(str(header_cell.value))

                # Check data rows
                for row_idx in range(5, 5 + len(export_data)):
                    cell = ws.cell(row=row_idx, column=col_idx)
                    if cell.value:
                        try:
                            cell_length = len(str(cell.value))
                            if cell_length > max_length:
                                max_length = cell_length
                        except Exception:
                            pass

                adjusted_width = min(max_length + 2, 50)
                ws.column_dimensions[column_letter].width = adjusted_width

            # Save to bytes
            output = io.BytesIO()
            wb.save(output)
            output.seek(0)

            return StreamingResponse(
                output,
                media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                headers={"Content-Disposition": f"attachment; filename={filename}.xlsx"},
            )

        else:  # CSV format
            csv_text = io.StringIO()
            if export_data:
                # Write header
                headers = list(export_data[0].keys())
                csv_text.write(','.join(headers) + '\n')

                # Write rows with special handling for phone numbers
                for row in export_data:
                    row_values = []
                    for header in headers:
                        value = row[header]
                        # Convert value to string
                        value_str = str(value)

                        # Special handling for Phone column - don't quote the formula
                        if header == 'Phone' and value_str.startswith('="'):
                            # Don't add extra quotes for formula cells
                            row_values.append(value_str)
                        elif ',' in value_str or '"' in value_str or '\n' in value_str:
                            # Quote fields that contain commas, quotes, or newlines
                            value_str = value_str.replace('"', '""')
                            row_values.append(f'"{value_str}"')
                        else:
                            row_values.append(value_str)

                    csv_text.write(','.join(row_values) + '\n')

            # Convert to bytes
            csv_bytes = io.BytesIO(csv_text.getvalue().encode('utf-8'))

            return StreamingResponse(
                csv_bytes,
                media_type="text/csv",
                headers={"Content-Disposition": f"attachment; filename={filename}.csv"},
            )

    except HTTPException:
        # Re-raise HTTP exceptions (like 404 for no data)
        raise
    except Exception as e:
        import traceback

        error_trace = traceback.format_exc()
        logger.error(f"  Error exporting candidates: {str(e)}")
        logger.error(error_trace)

        # Provide more detailed error message
        error_detail = f"Error exporting candidates: {str(e)}"
        if "openpyxl" in str(e).lower():
            error_detail += " (XLSX formatting error)"

        raise HTTPException(status_code=500, detail=error_detail)
